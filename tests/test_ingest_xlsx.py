"""Tests for the XLSX parser using a synthetic workbook generated in-test."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from ism_mcp.ingest import parse_xlsx


def _write_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controls - January 2026"
    ws.append(["ISM Controls"] + [""] * 24)
    ws.append([
        "Guideline", "Section", "Topic", "Identifier", "Revision", "Updated",
        "NC", "OS", "P", "S", "TS",
        "ML1", "ML2", "ML3",
        "Description",
        "Scoping1", "Scoping2", "Scoping3",
        "Provider", "Implementation Status", "Comments",
        "Consumer Responsibility", "Consumer Impl", "Consumer Config", "Comments",
    ])
    ws.append([
        "Guidelines for testing", "Encryption", "Network encryption",
        "ISM-9001", "1", "Jan-26",
        "Yes", "Yes", "Yes", "No", "No",
        "Yes", "No", "No",
        "All data communicated over network infrastructure is encrypted.",
        "", "", "", "", "Not Assessed", "", "", "Not Assessed", "Not Assessed", "",
    ])
    ws.append([
        "Guidelines for testing", "Audit", "Event logging",
        "ISM-9002", "2", "Feb-26",
        "No", "No", "Yes", "Yes", "Yes",
        "No", "Yes", "Yes",
        "Events are logged to a centralised facility.",
        "", "", "", "", "Not Assessed", "", "", "Not Assessed", "Not Assessed", "",
    ])
    wb.save(path)


def test_parse_extracts_all_ism_rows(tmp_path):
    workbook = tmp_path / "ccm.xlsx"
    _write_workbook(workbook)
    controls = list(parse_xlsx(workbook))
    assert {c.identifier for c in controls} == {"ISM-9001", "ISM-9002"}


def test_parse_populates_classification_and_maturity(tmp_path):
    workbook = tmp_path / "ccm.xlsx"
    _write_workbook(workbook)
    by_id = {c.identifier: c for c in parse_xlsx(workbook)}
    assert by_id["ISM-9001"].applies == {"NC": True, "OS": True, "P": True, "S": False, "TS": False}
    assert by_id["ISM-9001"].maturity == {"ML1": True, "ML2": False, "ML3": False}
    assert by_id["ISM-9002"].applies == {"NC": False, "OS": False, "P": True, "S": True, "TS": True}
    assert by_id["ISM-9002"].maturity == {"ML1": False, "ML2": True, "ML3": True}


def test_parse_skips_non_ism_rows(tmp_path):
    workbook = tmp_path / "ccm.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Controls - January 2026"
    ws.append(["ISM Controls"] + [""] * 24)
    ws.append([
        "Guideline", "Section", "Topic", "Identifier", "Revision", "Updated",
        "NC", "OS", "P", "S", "TS",
        "ML1", "ML2", "ML3",
        "Description",
        "x", "x", "x", "x", "x", "x", "x", "x", "x", "x",
    ])
    ws.append(["g", "s", "t", "not-an-id", "1", "Jan-26", "Yes", "Yes", "Yes", "No", "No",
               "No", "No", "No", "desc", "", "", "", "", "", "", "", "", "", ""])
    wb.save(workbook)
    assert list(parse_xlsx(workbook)) == []


def test_parse_raises_when_controls_sheet_missing(tmp_path):
    workbook = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "OnlyInfo"
    wb.save(workbook)
    with pytest.raises(ValueError, match="no 'Controls - ...' sheet"):
        list(parse_xlsx(workbook))
