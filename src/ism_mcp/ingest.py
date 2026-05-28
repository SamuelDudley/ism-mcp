"""Ingest the ASD ISM Cloud Controls Matrix XLSX and (optionally) the ISM PDF."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterator

import openpyxl
import pdfplumber

from .store import CLASSIFICATIONS, MATURITIES, Control


CONTROLS_SHEET = "Controls - March 2026"
HEADER_ROW = 2
DATA_START_ROW = 3


def parse_xlsx(xlsx_path: Path) -> Iterator[Control]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    # Sheet name embeds the revision date. Match by prefix so the parser
    # survives quarterly revisions without code changes.
    sheet_name = next(
        (sn for sn in wb.sheetnames if sn.startswith("Controls -")),
        None,
    )
    if sheet_name is None:
        raise ValueError(f"no 'Controls - ...' sheet found in {xlsx_path}")
    ws = wb[sheet_name]

    headers = [c.value for c in next(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW))]
    col = {name: i for i, name in enumerate(headers) if name}

    required = {"Identifier", "Description", "Guideline", "Section", "Topic"}
    missing = required - set(col)
    if missing:
        raise ValueError(f"XLSX header missing required columns: {missing}")

    for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        identifier = row[col["Identifier"]]
        if not identifier or not str(identifier).startswith("ISM-"):
            continue
        yield Control(
            identifier=str(identifier).strip(),
            guideline=str(row[col["Guideline"]] or "").strip(),
            section=str(row[col["Section"]] or "").strip(),
            topic=str(row[col["Topic"]] or "").strip(),
            revision=_str_or_none(row[col.get("Revision")]) if "Revision" in col else None,
            updated=_str_or_none(row[col.get("Updated")]) if "Updated" in col else None,
            description=str(row[col["Description"]] or "").strip(),
            applies={c: _yes(row[col[c]]) if c in col else False for c in CLASSIFICATIONS},
            maturity={m: _yes(row[col[m]]) if m in col else False for m in MATURITIES},
        )


def attach_pdf_excerpts(controls: list[Control], pdf_path: Path) -> list[Control]:
    """Walk the PDF once, find each ISM identifier occurrence, attach a paragraph excerpt and page number."""
    excerpts: dict[str, tuple[str, int]] = {}
    id_re = re.compile(r"\b(ISM-\d{3,4})\b")

    with pdfplumber.open(pdf_path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if "ISM-" not in text:
                continue
            for paragraph in _paragraphs(text):
                ids_in_para = set(id_re.findall(paragraph))
                if not ids_in_para:
                    continue
                for cid in ids_in_para:
                    if cid not in excerpts:
                        excerpts[cid] = (paragraph.strip(), page_no)

    return [
        Control(
            **{**c.as_dict(), "pdf_excerpt": excerpts.get(c.identifier, (None, None))[0],
               "pdf_page": excerpts.get(c.identifier, (None, None))[1]}
        )
        if c.identifier in excerpts
        else c
        for c in controls
    ]


def _paragraphs(text: str) -> Iterator[str]:
    para: list[str] = []
    for line in text.splitlines():
        if line.strip():
            para.append(line)
        elif para:
            yield " ".join(para)
            para = []
    if para:
        yield " ".join(para)


def _str_or_none(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _yes(v) -> bool:
    return str(v).strip().lower() == "yes"
